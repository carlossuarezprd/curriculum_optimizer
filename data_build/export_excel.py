"""Export the consolidated DB to a single Excel workbook (build/curriculum.xlsx).

Main sheet `sections` is the section-grain relational table from CLAUDE.md:
one row per section, course-level fields repeated, one boolean column per
concentration. Supporting sheets hold the items CLAUDE.md says to store
separately (concentration thresholds, core areas, flagship config) + full
course descriptions.

Regenerate:  python -m data_build.export_excel
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

OUT = Path("build/curriculum.xlsx")


def _autosize(ws, max_w=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max_w, max(10, width + 2))


def _header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"


def export():
    b = json.load(open("build/curriculum.json"))
    conc_names = [c["name"] for c in b["concentrations"]]
    wb = Workbook()

    # ---- sheet 1: sections (the consolidated DB) ----
    ws = wb.active
    ws.title = "sections"
    cols = ["course_number", "course_name", "units", "section_code", "quarter", "year",
            "professor", "time", "strict_prereqs", "recommended_prereqs", "core_area",
            "independent_application_course", "flagship_course",
            "r1_price_new", "r1_price_returning", "bid_matched"] + conc_names
    ws.append(cols)
    for s in b["sections"]:
        sc = set(s["concentrations"])
        ws.append([
            s["course_number"], s["course_name"], s["units"], s["section_code"],
            s["quarter"], s["year"], s["professor"], s["time"],
            " ".join(s["strict_prereqs"]), " ".join(s["recommended_prereqs"]),
            s["core_area"], bool(s["independent_application_course"]),
            bool(s["flagship_course"]), s["r1_price_new"], s["r1_price_returning"],
            bool(s["bid_matched"]),
        ] + [(n in sc) for n in conc_names])
    _header(ws)
    _autosize(ws)
    ws.column_dimensions["B"].width = 45  # course_name

    # ---- sheet 2: concentrations (thresholds stored separately) ----
    ws = wb.create_sheet("concentrations")
    ws.append(["concentration", "total_units_required", "n_qualifying_courses",
               "sub_buckets", "special", "notes"])
    for c in b["concentrations"]:
        buckets = " | ".join(f"{bk['label'] or 'list'}: {bk['units_required']}u"
                             + (f" (cap {bk['cap']})" if bk["cap"] else "")
                             for bk in c["buckets"])
        ws.append([c["name"], c["total_units_required"], len(c["qualifying"]),
                   buckets, bool(c["special"]), " ".join(c["notes"])])
    _header(ws); _autosize(ws)

    # ---- sheet 3: core_areas ----
    ws = wb.create_sheet("core_areas")
    ws.append(["area", "type", "units_required_each", "n_qualifying", "basic", "substitutes"])
    for a in b["core_areas"]:
        ws.append([a["name"], a["type"], b["meta"]["units_per_area"], len(a["qualifying"]),
                   " ".join(a["basic"]), " ".join(a["substitutes"])])
    _header(ws); _autosize(ws)

    # ---- sheet 4: flagship_config ----
    ws = wb.create_sheet("flagship_config")
    ws.append(["course_number", "professor", "label"])
    cname = {c["course_number"]: c["course_name"] for c in b["courses"]}
    for fp in b["flagship_seed"]:
        ws.append([fp["course_number"], fp["professor"], fp["label"]])
    _header(ws); _autosize(ws)

    # ---- sheet 5: courses (full descriptions) ----
    ws = wb.create_sheet("courses")
    ws.append(["course_number", "course_name", "units", "core_area", "concentrations",
               "application", "strict_prereqs", "recommended_prereqs", "description"])
    for c in b["courses"]:
        ws.append([c["course_number"], c["course_name"], c["units"], c["core_area"],
                   "; ".join(c["concentrations"]),
                   bool(c["independent_application_course"]),
                   " ".join(c["strict_prereqs"]), " ".join(c["recommended_prereqs"]),
                   c["description"]])
    _header(ws); _autosize(ws)
    ws.column_dimensions["I"].width = 80
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    # ---- sheet 6: meta / rules ----
    ws = wb.create_sheet("meta")
    ws.append(["key", "value"])
    for k, v in b["meta"].items():
        ws.append([k, v])
    _header(ws); _autosize(ws)

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    export()
